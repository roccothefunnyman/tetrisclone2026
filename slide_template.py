"""
AHEAD Workshop Slide Template Configuration

This template defines the styling and structure for generating workshop slides.
To create a new slide, use the create_slide_from_template() function with your content.

Example usage:
    from slide_template import create_slide_from_template

    content = {
        'title': 'Workshop 6: My New Workshop',
        'timing': 'Full Day | Week 6',
        'objective': 'Your objective text here.',
        'focus_items': [
            '- Focus item 1',
            '- Focus item 2',
            '- Focus item 3'
        ],
        'attendees': [
            '- Attendee type 1',
            '- Attendee type 2',
            '- Attendee type 3'
        ],
        'agenda_items': [
            'First agenda item',
            'Second agenda item',
            'Third agenda item'
        ],
        'flow_row1': [
            {'text': 'STEP1', 'subtitle': 'Description 1', 'color': 'teal'},
            {'text': 'STEP2', 'subtitle': 'Description 2', 'color': 'teal'},
            {'text': 'STEP3', 'subtitle': 'Description 3', 'color': 'teal'}
        ],
        'flow_row2': [
            {'text': 'STEP4', 'subtitle': 'Description 4', 'color': 'navy'},
            {'text': 'STEP5', 'subtitle': 'Description 5', 'color': 'navy'},
            {'text': 'STEP6', 'subtitle': 'Description 6', 'color': 'navy'}
        ],
        'outcome': 'Your key outcome text here.',
        'page_num': '6'
    }

    create_slide_from_template(content, 'my_new_slide.pptx')
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR


# =============================================================================
# BRAND COLORS - AHEAD Official Colors
# =============================================================================
COLORS = {
    'dark_navy': RGBColor(0x0D, 0x21, 0x37),    # #0D2137 - Primary dark color
    'teal': RGBColor(0x00, 0xA9, 0x9D),          # #00A99D - Accent color
    'orange': RGBColor(0xF7, 0x94, 0x1D),        # #F7941D - Highlight color
    'white': RGBColor(0xFF, 0xFF, 0xFF),         # #FFFFFF
    'light_gray': RGBColor(0xF0, 0xF0, 0xF0),    # #F0F0F0 - Background
    'light_teal_bg': RGBColor(0xE8, 0xF7, 0xF6), # #E8F7F6 - Key Outcome background
    'gray_text': RGBColor(0x66, 0x66, 0x66),     # #666666 - Subtitle text
}

# =============================================================================
# TYPOGRAPHY
# =============================================================================
FONT = {
    'family': 'Tenorite',  # Primary font for all text
    'sizes': {
        'title': Pt(28),           # Main slide title
        'section_header': Pt(16),  # Section titles (Agenda, Target Attendees, etc.)
        'body': Pt(14),            # Regular body text
        'body_small': Pt(12),      # Smaller body text (attendees list)
        'agenda_item': Pt(13),     # Agenda numbered items
        'label': Pt(12),           # Labels like "OBJECTIVE"
        'flow_box': Pt(14),        # Text inside flow boxes
        'subtitle': Pt(10),        # Flow box subtitles
        'footer': Pt(14),          # Footer tagline
        'page_num': Pt(12),        # Page number
    }
}

# =============================================================================
# SLIDE DIMENSIONS (16:9 Widescreen)
# =============================================================================
SLIDE = {
    'width': Inches(16.67),
    'height': Inches(9.375),
    'canvas_width': 1600,   # Reference canvas width in pixels
    'canvas_height': 900,   # Reference canvas height in pixels
}

# =============================================================================
# BOX STYLING
# =============================================================================
BOX_STYLE = {
    'fill_color': COLORS['white'],
    'border_color': COLORS['dark_navy'],
    'border_width': Pt(2),
    'corner_radius': 0.05,  # Rounded corner adjustment
}

# =============================================================================
# LAYOUT POSITIONS (in canvas pixels - will be scaled)
# =============================================================================
LAYOUT = {
    'header': {
        'height': 80,
    },
    'timing_badge': {
        'left': 1380, 'top': 20, 'width': 190, 'height': 40,
    },
    'objective': {
        'left': 30, 'top': 100, 'width': 1540, 'height': 90,
    },
    'focus': {
        'left': 30, 'top': 210, 'width': 380, 'height': 240,
    },
    'attendees': {
        'left': 30, 'top': 470, 'width': 380, 'height': 320,
    },
    'agenda': {
        'left': 430, 'top': 210, 'width': 580, 'height': 580,
    },
    'workshop_flow': {
        'left': 1030, 'top': 210, 'width': 540, 'height': 580,
    },
    'key_outcome': {
        'left': 1060, 'top': 540, 'width': 480, 'height': 220,
    },
    'footer': {
        'top': 850, 'height': 50,
    },
    'flow_box': {
        'width': 130, 'height': 50, 'spacing': 165,
        'row1_top': 300, 'row2_top': 420,
        'start_x': 1060,
    },
}


# =============================================================================
# HELPER FUNCTIONS
# =============================================================================

def scale_x(val):
    """Scale X coordinate from canvas pixels to slide inches."""
    return Inches(val * 16.67 / SLIDE['canvas_width'])


def scale_y(val):
    """Scale Y coordinate from canvas pixels to slide inches."""
    return Inches(val * 9.375 / SLIDE['canvas_height'])


def add_rectangle(slide, left, top, width, height, fill_color=None, line_color=None, line_width=Pt(0)):
    """Add a rounded rectangle shape."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    if fill_color:
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()

    if line_color:
        shape.line.color.rgb = line_color
        shape.line.width = line_width
    else:
        shape.line.fill.background()

    shape.adjustments[0] = BOX_STYLE['corner_radius']
    return shape


def add_plain_rectangle(slide, left, top, width, height, fill_color=None):
    """Add a plain rectangle (no rounded corners)."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    if fill_color:
        shape.fill.fore_color.rgb = fill_color
    shape.line.fill.background()
    return shape


def add_text_box(slide, left, top, width, height, text, font_size=Pt(16),
                 font_color=None, bold=False, align=PP_ALIGN.LEFT):
    """Add a text box with specified formatting."""
    if font_color is None:
        font_color = COLORS['dark_navy']

    textbox = slide.shapes.add_textbox(left, top, width, height)
    tf = textbox.text_frame
    tf.word_wrap = True
    tf.auto_size = None

    p = tf.paragraphs[0]
    p.text = text
    p.font.size = font_size
    p.font.color.rgb = font_color
    p.font.bold = bold
    p.font.name = FONT['family']
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(0)

    return textbox


def add_arrow(slide, left, top, width, height):
    """Add an arrow shape."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLORS['orange']
    shape.line.fill.background()
    return shape


# =============================================================================
# MAIN SLIDE CREATION FUNCTION
# =============================================================================

def create_slide_from_template(content, output_path):
    """
    Create a PowerPoint slide from the provided content dictionary.

    Args:
        content: Dictionary with keys:
            - title: str - Main slide title
            - timing: str - Timing badge text (e.g., "Half Day | Week 1")
            - objective: str - Objective description
            - focus_items: list[str] - AI Operating Model Focus bullet points
            - attendees: list[str] - Target Attendees bullet points
            - agenda_items: list[str] - Agenda items (numbers added automatically)
            - flow_row1: list[dict] - Top row flow boxes with 'text', 'subtitle', 'color'
            - flow_row2: list[dict] - Bottom row flow boxes with 'text', 'subtitle', 'color'
            - outcome: str - Key Outcome text
            - page_num: str - Page number for footer
        output_path: str - Path for the output .pptx file

    Returns:
        str - Path to the created file
    """
    prs = Presentation()
    prs.slide_width = SLIDE['width']
    prs.slide_height = SLIDE['height']

    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # Background
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = COLORS['light_gray']

    # Header Bar
    add_plain_rectangle(slide, Inches(0), Inches(0), SLIDE['width'],
                        scale_y(LAYOUT['header']['height']), COLORS['dark_navy'])

    # Title
    add_text_box(slide, scale_x(40), scale_y(15), scale_x(900), scale_y(50),
                 content['title'], font_size=FONT['sizes']['title'], font_color=COLORS['white'])

    # Timing Badge
    L = LAYOUT['timing_badge']
    badge = add_rectangle(slide, scale_x(L['left']), scale_y(L['top']),
                          scale_x(L['width']), scale_y(L['height']), fill_color=COLORS['orange'])
    tf = badge.text_frame
    p = tf.paragraphs[0]
    p.text = content['timing']
    p.font.size = Pt(14)
    p.font.color.rgb = COLORS['white']
    p.font.bold = True
    p.font.name = FONT['family']
    p.alignment = PP_ALIGN.CENTER

    # Objective Section
    L = LAYOUT['objective']
    add_rectangle(slide, scale_x(L['left']), scale_y(L['top']), scale_x(L['width']), scale_y(L['height']),
                  fill_color=BOX_STYLE['fill_color'], line_color=BOX_STYLE['border_color'],
                  line_width=BOX_STYLE['border_width'])
    obj_box = slide.shapes.add_textbox(scale_x(50), scale_y(115), scale_x(1480), scale_y(60))
    tf = obj_box.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.text = "OBJECTIVE"
    p1.font.size = FONT['sizes']['label']
    p1.font.color.rgb = COLORS['dark_navy']
    p1.font.bold = True
    p1.font.name = FONT['family']
    p2 = tf.add_paragraph()
    p2.text = content['objective']
    p2.font.size = FONT['sizes']['section_header']
    p2.font.color.rgb = COLORS['dark_navy']
    p2.font.name = FONT['family']
    p2.space_before = Pt(8)

    # AI Operating Model Focus
    L = LAYOUT['focus']
    add_rectangle(slide, scale_x(L['left']), scale_y(L['top']), scale_x(L['width']), scale_y(L['height']),
                  fill_color=BOX_STYLE['fill_color'], line_color=BOX_STYLE['border_color'],
                  line_width=BOX_STYLE['border_width'])
    focus_box = slide.shapes.add_textbox(scale_x(50), scale_y(225), scale_x(340), scale_y(210))
    tf = focus_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "AI Operating Model Focus"
    p.font.size = FONT['sizes']['section_header']
    p.font.color.rgb = COLORS['dark_navy']
    p.font.bold = True
    p.font.name = FONT['family']
    for item in content['focus_items']:
        p = tf.add_paragraph()
        p.text = item
        p.font.size = FONT['sizes']['body']
        p.font.color.rgb = COLORS['dark_navy']
        p.font.name = FONT['family']
        p.space_before = Pt(12)

    # Target Attendees
    L = LAYOUT['attendees']
    add_rectangle(slide, scale_x(L['left']), scale_y(L['top']), scale_x(L['width']), scale_y(L['height']),
                  fill_color=BOX_STYLE['fill_color'], line_color=BOX_STYLE['border_color'],
                  line_width=BOX_STYLE['border_width'])
    att_box = slide.shapes.add_textbox(scale_x(50), scale_y(485), scale_x(350), scale_y(290))
    tf = att_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Target Attendees"
    p.font.size = FONT['sizes']['section_header']
    p.font.color.rgb = COLORS['dark_navy']
    p.font.bold = True
    p.font.name = FONT['family']
    for item in content['attendees']:
        p = tf.add_paragraph()
        p.text = item
        p.font.size = FONT['sizes']['body_small']
        p.font.color.rgb = COLORS['dark_navy']
        p.font.name = FONT['family']
        p.space_before = Pt(10)

    # Agenda
    L = LAYOUT['agenda']
    add_rectangle(slide, scale_x(L['left']), scale_y(L['top']), scale_x(L['width']), scale_y(L['height']),
                  fill_color=BOX_STYLE['fill_color'], line_color=BOX_STYLE['border_color'],
                  line_width=BOX_STYLE['border_width'])
    agenda_box = slide.shapes.add_textbox(scale_x(455), scale_y(230), scale_x(540), scale_y(540))
    tf = agenda_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Agenda"
    p.font.size = FONT['sizes']['section_header']
    p.font.color.rgb = COLORS['dark_navy']
    p.font.bold = True
    p.font.name = FONT['family']
    for i, item in enumerate(content['agenda_items'], 1):
        p = tf.add_paragraph()
        p.text = f"{i}.  {item}"
        p.font.size = FONT['sizes']['agenda_item']
        p.font.color.rgb = COLORS['dark_navy']
        p.font.name = FONT['family']
        p.space_before = Pt(14)

    # Workshop Flow
    L = LAYOUT['workshop_flow']
    add_rectangle(slide, scale_x(L['left']), scale_y(L['top']), scale_x(L['width']), scale_y(L['height']),
                  fill_color=BOX_STYLE['fill_color'], line_color=BOX_STYLE['border_color'],
                  line_width=BOX_STYLE['border_width'])
    add_text_box(slide, scale_x(1030), scale_y(230), scale_x(540), scale_y(35),
                 "Workshop Flow", font_size=FONT['sizes']['section_header'],
                 font_color=COLORS['dark_navy'], bold=True, align=PP_ALIGN.CENTER)

    # Flow Row 1
    FL = LAYOUT['flow_box']
    x_pos = FL['start_x']
    for i, flow in enumerate(content['flow_row1']):
        color = COLORS['teal'] if flow.get('color', 'teal') == 'teal' else COLORS['dark_navy']
        box = add_rectangle(slide, scale_x(x_pos), scale_y(FL['row1_top']),
                            scale_x(FL['width']), scale_y(FL['height']), fill_color=color)
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = flow['text']
        p.font.size = FONT['sizes']['flow_box']
        p.font.color.rgb = COLORS['white']
        p.font.bold = True
        p.font.name = FONT['family']
        p.alignment = PP_ALIGN.CENTER
        if flow.get('subtitle'):
            add_text_box(slide, scale_x(x_pos), scale_y(355), scale_x(FL['width']), scale_y(25),
                         flow['subtitle'], font_size=FONT['sizes']['subtitle'],
                         font_color=COLORS['gray_text'], align=PP_ALIGN.CENTER)
        if i < len(content['flow_row1']) - 1:
            add_arrow(slide, scale_x(x_pos + 138), scale_y(318), scale_x(30), scale_y(15))
        x_pos += FL['spacing']

    # Flow Row 2
    x_pos = FL['start_x']
    for i, flow in enumerate(content['flow_row2']):
        color = COLORS['dark_navy'] if flow.get('color', 'navy') == 'navy' else COLORS['teal']
        box = add_rectangle(slide, scale_x(x_pos), scale_y(FL['row2_top']),
                            scale_x(FL['width']), scale_y(FL['height']), fill_color=color)
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = flow['text']
        p.font.size = FONT['sizes']['flow_box']
        p.font.color.rgb = COLORS['white']
        p.font.bold = True
        p.font.name = FONT['family']
        p.alignment = PP_ALIGN.CENTER
        if flow.get('subtitle'):
            add_text_box(slide, scale_x(x_pos), scale_y(475), scale_x(FL['width']), scale_y(25),
                         flow['subtitle'], font_size=FONT['sizes']['subtitle'],
                         font_color=COLORS['gray_text'], align=PP_ALIGN.CENTER)
        if i < len(content['flow_row2']) - 1:
            add_arrow(slide, scale_x(x_pos + 138), scale_y(438), scale_x(30), scale_y(15))
        x_pos += FL['spacing']

    # Key Outcome
    L = LAYOUT['key_outcome']
    add_rectangle(slide, scale_x(L['left']), scale_y(L['top']), scale_x(L['width']), scale_y(L['height']),
                  fill_color=COLORS['light_teal_bg'], line_color=BOX_STYLE['border_color'],
                  line_width=BOX_STYLE['border_width'])
    outcome_box = slide.shapes.add_textbox(scale_x(1060), scale_y(560), scale_x(480), scale_y(180))
    tf = outcome_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Key Outcome"
    p.font.size = FONT['sizes']['body']
    p.font.color.rgb = COLORS['dark_navy']
    p.font.bold = True
    p.font.name = FONT['family']
    p.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph()
    p2.text = content['outcome']
    p2.font.size = FONT['sizes']['body_small']
    p2.font.color.rgb = COLORS['dark_navy']
    p2.font.name = FONT['family']
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(16)

    # Footer
    add_plain_rectangle(slide, Inches(0), scale_y(LAYOUT['footer']['top']),
                        SLIDE['width'], scale_y(LAYOUT['footer']['height']), COLORS['dark_navy'])
    add_text_box(slide, scale_x(30), scale_y(862), scale_x(40), scale_y(30),
                 content['page_num'], font_size=FONT['sizes']['page_num'], font_color=COLORS['white'])
    add_text_box(slide, scale_x(550), scale_y(862), scale_x(500), scale_y(30),
                 "We bring AI to life in the Enterprise", font_size=FONT['sizes']['footer'],
                 font_color=COLORS['white'], align=PP_ALIGN.CENTER)
    add_text_box(slide, scale_x(1420), scale_y(862), scale_x(150), scale_y(30),
                 "AHEAD", font_size=FONT['sizes']['section_header'],
                 font_color=COLORS['white'], bold=True, align=PP_ALIGN.RIGHT)

    prs.save(output_path)
    print(f"Created: {output_path}")
    return output_path


# =============================================================================
# EXAMPLE / TEST
# =============================================================================

if __name__ == "__main__":
    # Example: Create a test slide
    example_content = {
        'title': 'Workshop 6: Example Workshop Title',
        'timing': 'Half Day | Week 6',
        'objective': 'This is an example objective describing what this workshop will accomplish.',
        'focus_items': [
            '- First focus area',
            '- Second focus area',
            '- Third focus area'
        ],
        'attendees': [
            '- Executive sponsors',
            '- Technical leads',
            '- Project managers',
            '- Key stakeholders'
        ],
        'agenda_items': [
            'Introduction and overview',
            'Deep dive into topic A',
            'Hands-on exercise',
            'Discussion and Q&A',
            'Next steps and action items'
        ],
        'flow_row1': [
            {'text': 'DISCOVER', 'subtitle': 'Requirements', 'color': 'teal'},
            {'text': 'DESIGN', 'subtitle': 'Architecture', 'color': 'teal'},
            {'text': 'DEVELOP', 'subtitle': 'Implementation', 'color': 'teal'}
        ],
        'flow_row2': [
            {'text': 'TEST', 'subtitle': 'Validation', 'color': 'navy'},
            {'text': 'DEPLOY', 'subtitle': 'Release', 'color': 'navy'},
            {'text': 'MONITOR', 'subtitle': 'Operations', 'color': 'navy'}
        ],
        'outcome': 'Clear understanding of the process and defined next steps for implementation.',
        'page_num': '6'
    }

    create_slide_from_template(example_content, 'example_slide.pptx')
